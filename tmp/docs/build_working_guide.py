from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(r'C:/xampp/htdocs/lms')
OUT = ROOT / 'output/docx/TLMTI_Skills_LMS_Working_Guide_Hinglish.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.75)
s.header_distance = s.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = d.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
normal = d.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.06
for name, size in [('Title', 24), ('Heading 1', 18), ('Heading 2', 12)]:
    st = d.styles[name]
    st.font.size = Pt(size)
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.keep_with_next = True
d.styles['Subtitle'].font.size = Pt(11)
header = s.header.paragraphs[0]
header.text = 'TLMTI Skills LMS  |  Working guide aur review checklist'
header.style = d.styles['Normal']
header.runs[0].font.size = Pt(9)
footer = s.footer.paragraphs[0]
footer.text = 'Review copy  |  27 October 2026                                           Page '
footer.runs[0].font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
d.core_properties.title = 'TLMTI Skills LMS Working Guide'
d.core_properties.subject = 'Hinglish user workflows, PWA scope and acceptance checklist'
for style in d.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t, 2)
def b(t): d.add_paragraph(t, 'List Bullet')
def page(title):
    d.add_page_break()
    d.add_heading(title, 1)
def table(headers, rows, widths):
    t = d.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i, text in enumerate(headers): t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row): cells[i].text = str(text)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:'+side); el.set(qn('w:w'), '85'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            pr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el = OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            pr.append(borders)
            if ri == 0 or ri % 2 == 0:
                fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'DFEAF0' if ri == 0 else 'F5F8FA'); pr.append(fill)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1
                for run in para.runs:
                    run.font.size = Pt(10)
                    if ri == 0: run.bold = True
    d.add_paragraph().paragraph_format.space_after = Pt(0)

d.add_heading('TLMTI Skills LMS Working Guide', 0)
p('Hinglish guide  |  Local review version  |  27 October 2026')
p('Yeh guide project owner, tester aur training team ke liye hai. Isse aap samajh sakte hain ki LMS kaise chalega, har role kya karega aur local system ko kaise check karna hai. Current implementation aur future work alag bataye gaye hain; ise production completion certificate na samjhein.')
h('1 Project ka purpose aur agreed scope')
p('Admin training centres, users, subjects aur enrolments set karega. Teacher apne assigned subjects ka content, tests aur attendance sambhalega. Student enrolled courses padhega, tests dega aur apni progress dekhega. Data server par save hoga aur permission ke hisab se hi dikhega.')
p('Mobile delivery PWA hogi. Flutter app is scope mein nahi hai. Desktop aur mobile par wahi responsive web application chalegi. Installable PWA ki final readiness aur real phone verification abhi pending hain.')
table(['Hissa', 'Abhi ki sthiti'], [
('Admin aur Teacher', 'Core workflows implemented; owner acceptance checklist is guide mein hai.'),
('Student', 'Learning, tests, progress, attendance aur selected offline downloads implemented.'),
('PWA', 'Manifest aur service worker foundation hai; install wiring aur device acceptance pending.'),
('Redis', 'Abhi file sessions. Redis session support code hai; live activation aur verification pending.'),
('Storage', 'Local uploads. Cloudflare R2 baad mein, owner ke kehne par.'),
('Theme aur language', 'Light only; Hindi aur English. Chhattisgarhi selectable nahi hai.')
], [1.55, 5.45])
h('Guide ka reading order')
p('Pages 2–3 Admin setup; page 4 Teacher; page 5 Student; page 6 learning rules; page 7 PWA aur offline; pages 8–9 acceptance tests; page 10 technical handover aur pending work.')
p('Review ke baad next instruction milne tak development, Redis activation aur PWA changes par kaam rokna hai. Is document ko banate samay application code ya database badalne ka scope nahi hai.')

page('2 Admin se system setup kaise hoga')
p('Local login URL: http://localhost/lms/public/login.php. Pehle Apache aur MySQL running hon. Existing database aur project configuration preserve karein; demo data dobara insert karna review ke liye zaroori nahi hai.')
h('Setup ka sahi order')
b('Centres: centre name, district, address aur active status save karein. Batches: centre, residential/community mode aur valid start/end dates set karein.')
b('Users: teacher aur student accounts banayein; role, centre, contact aur status sahi choose karein. Password hashed save hota hai. Admin account password reset kar sakta hai.')
b('Subjects: unique subject/trade code, name, description aur active status save karein. Duplicate code allow nahi hona chahiye.')
b('Teacher assignment: Subjects ke assignments tab mein teacher ko subject dein. Content manage aur test manage flags alag hain; zaroorat ke hisab se enable karein.')
b('Enrolments: student ko subject aur optional batch se link karein. Centre consistent ho. Active enrolment se learning start hogi aur week one khulega.')
b('Content aur tests publish karein; phir student account se course visibility aur unlocking verify karein. Unpublished records student ko learning material ke roop mein nahi dikhne chahiye.')
h('Users aur record status')
p('Admin apna account deactivate/demote nahi kar sakta; last active admin ko remove karne se protection hai. Enrolment ko completed ya dropped karne se history erase nahi hoti. Completed course review mode mein rahega; dropped enrolment se nayi learning aur attempts unavailable honge.')
h('Attendance aur reports')
p('Attendance ek person aur ek date ke liye save hoti hai. Present, absent, late, half day, leave ya holiday aur remarks use karein. Same person/date duplicate reject hoga; galti ko existing entry edit karke correct karein. Yeh per subject class attendance model nahi hai.')
p('Reports mein subject filters, progress, test averages, attendance counts aur student result details milte hain. CSV export matching records deta hai. Screen par list pagination hone ka matlab CSV sirf current page export kare, aisa nahi hai.')
h('Admin ke additional controls')
p('Roles/permissions, application settings, audit logs, profile/password aur personal notifications available hain. Security aur test settings badalne ke baad affected flow ko phir check karein. Existing reports ya users ko dummy numbers se replace nahi karna hai.')

page('3 Content aur test authoring')
h('Modules aur lessons')
p('Content page par module ka subject, week, sequence, estimated minutes aur publication status set karein. Hindi aur English titles/descriptions author ko khud bharne hain. UI language selector content ka automatic translation nahi karta.')
p('Lesson ko module se link karein. Video, audio, PDF, image, text, checklist aur step guide types available hain. Lesson sequence, mandatory flag, downloadable flag, minimum completion percentage aur bilingual body/title check karein. Published module ke andar published lessons student flow mein aate hain.')
h('Media upload')
p('Media tab se file upload karke lesson mein attach karein. Supported formats mein JPG, PNG, WebP, PDF, MP4, MP3, OGG aur WAV hain. Files local storage mein rehti hain; access login aur subject permission se protected hai.')
p('Video/audio ki duration confirm karein; supported browser mein duration detect ho sakti hai. Allowed input 1 se 86400 seconds hai. App limit 50 MB hai, lekin PHP upload/post limit isse kam ho sakti hai. Upload page par effective server limit check karein.')
h('Question bank aur test')
b('Question bank mein subject/module, bilingual question, four options, ek correct answer, marks, difficulty aur status save karein. Current authoring UI single-answer MCQ ke liye hai.')
b('Module, weekly ya practice test choose karein. Title, subject/week/module, duration, pass percentage, pass comparison, attempt limit, retake wait, minimum study requirement aur publication set karein.')
b('Duration 0 ka matlab untimed; blank maximum attempts ka matlab unlimited. Pass comparison gte mein threshold ke barabar marks pass hain; gt mein usse zyada chahiye.')
b('Usi subject ke published questions select karein. Total marks selected questions se calculate hote hain. Student account se test availability aur result check karein.')
h('Published assessments ko badalne ka rule')
p('Jis test ya question par attempt start ho chuka ho, pending attempt samet, usse in place edit nahi kiya ja sakta. Correction ke liye naya question/test record banayein. Isse purane attempts ka question, options, marks aur pass rule stable rehta hai.')
p('Example: module test ke liye 60% pass rule aur mandatory lesson completion rakhein. Student lesson complete kare, test de aur pass hone par next module dekhe. Sirf test publish karne se locked learning sequence bypass nahi hota.')

page('4 Teacher ka daily workflow')
p('Teacher login ke baad sirf assigned subjects aur unmein actively enrolled students dekhega. Dashboard se subjects, student counts, scores aur recent activity tak ja sakta hai. Yeh scope reports, media, direct record links aur CSV exports par bhi apply hota hai.')
h('Class se pehle')
b('My Subjects mein subject kholein. Expected students na milne par Admin se subject assignment aur active enrolment verify karwayein.')
b('Content permission ho to modules/lessons aur media create/edit karein. Hindi aur English content alag verify karke publish karein.')
b('Test permission ho to question bank aur tests banayein. Sahi subject, pass rule, duration aur retry settings check karein.')
h('Class ke dauran aur baad')
b('Assigned students ki daily attendance mark karein. Attendance person/date basis par hai; same din alag subject ke liye duplicate row na banayein.')
b('Students/detail view se progress aur submitted test results dekhein. Low score ya incomplete study par student ko lesson review aur allowed retake ka direction dein.')
b('Reports ko subject ke hisab se filter karein aur zaroorat par CSV export karein. Student se credentials lekar uska account use karna teacher workflow nahi hai.')
h('Teacher permissions ki boundary')
table(['Action', 'Teacher access'], [
('Assigned subject dekhna', 'Haan, permitted scope ke andar.'),
('Content create ya edit', 'Role permission aur subject content flag dono required.'),
('Questions aur tests manage', 'Role permission aur subject test flag dono required.'),
('Attendance aur reports', 'Assigned student/subject scope tak.'),
('Users, centres, batches, enrolments', 'Admin management; Teacher create/edit nahi karega.'),
('Roles aur system settings', 'Admin management; Teacher ko access nahi.')
], [2.35, 4.65])
h('Teacher acceptance example')
p('demo.teacher.copa se login karein. Assigned COPA subject aur uske students dikhne chahiye. Kisi unassigned subject ka direct link open ya edit karne par access deny hona chahiye. Admin content/test flag disable kare to corresponding write action reject hona chahiye.')

page('5 Student ka end to end workflow')
h('Login aur dashboard')
p('Student apne account se login karega. Dashboard saved course progress, study time, completed lessons, attendance, continue learning aur next available test dikhata hai. Theory, practical ya placement percentages bina supporting records ke invent nahi hote.')
p('Hindi ya English select karein. Preference account/session mein save hoti hai. Interface light theme mein rahega. Desktop par sidebar aur mobile par menu/bottom navigation use honge; narrow tables apne card ke andar scroll kar sakti hain.')
h('Course se lesson tak')
b('My Courses kholein; search, status filter aur pagination se enrolled course choose karein. Modules/lessons ka completed, available ya locked state dekhein.')
b('Available lesson kholein. Text padhein, media play karein ya PDF dekhein. Video position save hoti hai aur available position se resume kiya ja sakta hai.')
b('Actual active study karein. Tab hidden ya playback pause hone par timer behavior badalta hai. Video ko aage seek karna skipped watch time earn nahi karta.')
b('Required study complete hone par Mark Complete karein. Sirf page open kar dena completion ke liye kaafi nahi hai. Uske baad available test ya next lesson kholein.')
h('Test dena')
p('Test start par server attempt create karta hai; dobara start par same pending attempt resume hota hai. Student ek question par answer choose karke next/previous kar sakta hai. Draft answers local browser mein account/attempt ke hisab se save hote hain.')
p('Timed test ka countdown dekhein aur deadline se pehle submit karein. Grading server karta hai. Submitted result mein score, pass/fail aur apna attempt history milega. Repeated submit se duplicate result ya extra marks nahi bante.')
h('Progress aur account')
p('My Progress se module progress aur apne result links dekhein. My Attendance read only hai; correction Teacher/Admin karega. Profile mein permitted name/contact changes aur current password verification ke saath password change karein. Personal notifications shared account page par milti hain.')
p('Completed enrolment ka published content review kiya ja sakta hai; naya progress/test write allowed nahi. Dropped enrolment ya inactive subject ki learning unavailable hogi, lekin apne submitted results history mein mil sakte hain.')

page('6 Learning aur assessment ke rules')
table(['Condition', 'System ka expected behavior'], [
('Naya active enrolment', 'Week one open; later weeks weekly pass se unlock.'),
('Ek week ke modules', 'Sequence mein khulenge. Mandatory lessons order mein complete honge.'),
('Optional lesson pending', 'Sequence ko block nahi karega.'),
('Module ke published tests', 'Required lessons aur sab published module tests pass hone par next module khulega.'),
('Weekly test pass', 'Agla week open hoga.'),
('Weekly test fail', 'Us week ka fresh study cycle; purana progress/result history preserve.'),
('Module test fail', 'Next module locked; configured retake rule se retry.'),
('Timed attempt late', 'Server deadline ke baad failed zero-score result; offline rehne se extra time nahi.'),
('Duplicate submission', 'Pehle se saved result return; duplicate answers nahi.'),
('Old cycle offline activity', 'Naye study cycle ki requirement satisfy nahi karegi.')
], [2.12, 4.88])
h('Study time ka calculation')
p('Video/audio ke liye required time minimum lesson seconds aur media duration × completion percentage mein se jo zyada ho, woh hota hai. Duration missing ho to configured minimum lesson seconds use hote hain. Online study seconds server ke session clock se cap hote hain.')
p('Example: 10 minute video aur 80% completion par 8 minute chahiye, agar configured minimum usse zyada na ho. Page kholkar ya seek karke 8 minute ka watch credit turant nahi milta.')
h('Test availability aur grading')
p('Module test mandatory lessons check karta hai. Weekly/practice test configured minimum study percentage aur availability window check karta hai; minimum 0 ho sakta hai. Attempts, wait interval, duration aur pass rule server enforce karta hai. Active student flow mein legacy 15 day module delay apply nahi hota.')
p('Correct answers student browser ko start payload mein nahi bheje jaate. Current authoring single-answer MCQ hai; runtime mein true/false support bhi hai. Advanced question types aur instructor grading ko complete feature na maanein.')

page('7 PWA aur offline ka working model')
p('PWA ka target ek hi web application ko browser aur supported install experience mein chalana hai. PHP backend, database aur role permissions wahi rahenge. App install hone se saara course content automatically offline nahi aata.')
h('Abhi offline kya available hai')
b('Online rehte hue downloadable lesson ko Downloads mein explicitly save karein. Saved lesson text aur fully downloaded media offline reopen ho sakte hain. Visited student pages bhi cache ho sakte hain.')
b('Study events aur test drafts IndexedDB mein user-wise store hote hain. App open hone par reconnect/resume ya Sync now se server sync try hota hai.')
b('Pehle se started assessment ka draft offline continue ho sakta hai; final grade ke liye server chahiye. Timed answers deadline ke andar server tak pahunchne chahiye.')
b('Logout/account switch par private page/media caches clear hote hain. Unsynced events user-wise partitioned rehte hain. Session expire ho to online login/page reload ke baad sync retry karein.')
h('Offline ki practical limits')
p('Unvisited ya undownloaded lessons offline guaranteed nahi hain. Browser storage eviction se downloads hat sakte hain. Closed app background sync guarantee aur push notifications implemented nahi hain. Admin/Teacher management writes online honge. Offline study duration independently verify nahi hoti; events offline mark hote hain aur per event four hours tak bounded hain.')
h('Installable PWA ke liye pending work')
p('Current project mein manifest.webmanifest aur registered service worker hain. Shared layout mein manifest link wiring missing mili; icon set abhi single SVG hai. Install icons/browser requirements aur install experience verify karna baaki hai. Isliye complete installable PWA ka signoff abhi nahi diya gaya hai.')
p('Next authorized phase mein manifest wiring, appropriate install icons, secure hosted test URL aur supported browser install/Add to Home Screen flow check hoga. Android/iPhone par standalone launch, navigation, offline download, logout isolation aur sync ki real device testing bhi pending hai.')
p('Desktop ka localhost PC ko point karta hai; phone ka localhost phone ko. Phone testing ke liye reachable test host aur suitable secure context chahiye. Is document review ke dauran network/server deployment setup change nahi kiya ja raha hai.')

page('8 Owner acceptance checklist')
p('Demo data synthetic hai; real student records nahi hain. Seed ke baad manual records judne se counts badal sakte hain. Tests ke liye demo accounts use karein aur apna observed result har row mein likhein. Password badla gaya ho to initial demo password kaam nahi karega.')
table(['Role', 'Local demo login', 'Password'], [
('Admin', 'admin', 'Owner ke paas existing password; shareable guide mein omit hai.'),
('Teacher', 'demo.teacher.copa', 'Initial demo password Demo@2026!'),
('Student', 'demo.student.01', 'Initial demo password Demo@2026!')
], [1.0, 2.3, 3.7])
p('Baaki teachers: demo.teacher.elec, demo.teacher.weld, demo.teacher.sew, demo.teacher.dsl. Students demo.student.01 se demo.student.24 tak. Demo accounts sirf local demonstration ke liye hain.')
table(['ID', 'Check aur expected result', 'Pass Fail aur note'], [
('A1', 'Admin login; centres, subjects aur users ki actual saved entries dikhengi.', ''),
('A2', 'Naya centre/batch save; invalid end date reject; edit refresh ke baad persist.', ''),
('A3', 'Demo teacher/student create ya edit; correct role/status/contact persist.', ''),
('A4', 'Subject create; duplicate code reject; teacher assignment save.', ''),
('A5', 'Student enrolment aur centre consistency; active student ko subject visible.', ''),
('A6', 'Bilingual module/lesson/media save; published content student ko visible.', ''),
('A7', 'MCQ/test save; total marks correct; attempted question/test edit blocked.', ''),
('A8', 'Attendance save/edit; same user/date duplicate reject.', ''),
('A9', 'Report filter aur CSV values saved data se match; pagination stable.', ''),
('T1', 'Teacher sirf assigned subjects/students aur allowed reports dekhe.', ''),
('T2', 'Unassigned direct link/edit/export denied; disabled content/test flag enforced.', '')
], [.4, 5.15, 1.45])

page('9 Student aur PWA acceptance checklist')
table(['ID', 'Check aur expected result', 'Pass Fail aur note'], [
('S1', 'Student login; course list sirf apne enrolments ki ho.', ''),
('S2', 'Available lesson study, complete aur refresh; progress persist.', ''),
('S3', 'Locked lesson/module direct URL se bypass na ho.', ''),
('S4', 'Test answer/submit; score correct; duplicate submit same result de.', ''),
('S5', 'Module pass next module khole; weekly pass next week khole.', ''),
('S6', 'Weekly fail naya cycle; old result history mein rahe.', ''),
('S7', 'Timed test expiry par failed zero; client clock change se extension nahi.', ''),
('S8', 'Student kisi aur ka result/media access na kar sake.', ''),
('U1', 'Hindi/English switch aur relogin par preference; light UI consistent.', ''),
('U2', 'Phone/tablet/desktop par menus, forms aur buttons usable; overflow contained.', ''),
('P1', 'Online lesson/media download; actual network off karke content reopen.', ''),
('P2', 'Offline progress ke baad reconnect aur Sync now; correct account mein save.', ''),
('P3', 'Logout/account switch; previous user ke private downloads visible na hon.', ''),
('P4', 'Install/Add to Home Screen aur standalone phone launch; pending PWA fixes ke baad.', '')
], [.4, 5.15, 1.45])
h('Ab tak recorded verification')
p('Previous local checks mein 186 assertions pass record hue: core 43, staff CRUD 17, staff HTTP 64, student rules 24, student HTTP 32 aur service worker 6. PHP parse aur JavaScript syntax checks bhi pass record hue. Is document pass mein suites dobara run nahi ki gayi hain.')
p('Browser demo mein lesson completion, 100% server-scored module result, next module unlock aur two lesson downloads verify hue. Service worker network-loss tests simulated runtime mein the. Physical phone installation, actual phone offline acceptance, load testing aur production signoff pending hain.')

page('10 Technical handover aur next decisions')
h('Current architecture')
p('Core PHP aur PDO MySQL/MariaDB application XAMPP par chalti hai. Existing page routes aur authentication preserve hain. Staff dispatch aur StudentService role-specific behavior handle karte hain. Local database name lms hai; example configuration ko actual configuration par blindly copy na karein.')
p('Login sessions abhi PHP files mein hain. SessionStore mein Redis session driver support hai, lekin Redis service/phpredis extension ki live setup aur verification pending hai. Current login rate limit database-backed hai; Redis cache/rate limiting implemented maana nahi ja sakta.')
p('Lists SQL LIMIT/OFFSET use karti hain; backend bounded page sizes aur deterministic ordering rakhta hai. Media storage local storage/uploads mein aur metadata/checksums media_assets mein hain. public/media.php authenticated scoped delivery aur byte ranges handle karta hai.')
h('Relevant project references')
b('docs/ADMIN_TEACHER.md aur docs/STUDENT_MODULE.md: detailed implementation notes aur verification scope.')
b('public/manifest.webmanifest, public/sw.js aur public/assets/js/app-ui.js: current PWA/offline foundation. app/Support/ui.php: shared layout.')
b('storage/qa: recorded UI screenshots. tests: automated checks. Existing production-like records ko acceptance testing ke liye delete na karein.')
h('Agla kaam sirf owner instruction ke baad')
table(['Priority', 'Pending scope'], [
('1 Owner review', 'Is guide ka workflow aur acceptance notes confirm; changes list banayein.'),
('2 PWA completion', 'Manifest/layout integration, install assets aur actual device acceptance.'),
('3 Redis', 'Service/extension setup, session activation, login/logout/expiry verification; cache scope separately confirm.'),
('Later R2', 'Cloudflare R2 storage integration; filhaal local upload model continue.'),
('Separate scope', 'Advanced question types, practical evidence, certificates/badges aur placement features if required.'),
('Before production', 'HTTPS hosting, backup/restore, security/configuration review aur operational acceptance.')
], [1.35, 5.65])
h('Review record')
p('Reviewed by: ____________________    Date: ____________________')
p('Required changes / approval notes: ______________________________________')
p('____________________________________________________________________')
p('Owner ki next instruction aane tak development paused rahega. PWA confirmed delivery direction hai; Redis aur R2 ko completed integration label nahi diya gaya hai.')

for para in d.paragraphs:
    for border in list(para._p.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
d.save(OUT)
print(OUT)
